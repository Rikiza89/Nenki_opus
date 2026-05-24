"""Multi-stage Excel import validation pipeline.

The pipeline:
  1. Pulls required fields (name, death date) from each row using the column mapping.
  2. Parses the death date (any of the supported Japanese / Gregorian forms).
  3. Applies business rules (no future dates, no pre-Meiji dates).
  4. Builds dynamic attributes, applying the user's column-remap selections.
  5. Detects duplicates against the DB *and* against rows seen earlier in the
     same file (so a doubled row inside one Excel sheet is reported, not
     silently inserted twice).

Errors are recorded with the raw row data so the UI can offer the user
inline correction.
"""

import datetime
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from memorial_app.core.japanese_date_parser import (
    parse_date,
    parse_date_from_columns,
    DateValidationError,
)
from memorial_app.database.db_manager import DatabaseManager
from memorial_app.importer.excel_importer import ColumnMapping


# Cells that begin with one of these characters could be interpreted as a
# formula by Excel and trigger unintended computation when the user opens the
# error-export file. We prefix them with an apostrophe to neutralize that.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _sanitize_for_excel(value: str) -> str:
    if not value:
        return value
    if value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


@dataclass
class ValidatedRow:
    row_index: int
    name: str
    death_date: str  # ISO format
    era_display: str  # Japanese era display
    attributes: dict[str, str]
    source_file_path: str
    # Set by the UI when the user resolves a duplicate at confirm time.
    # Values: None (default — insert as new), "overwrite" (update existing
    # row with id=duplicate_of), or "skip" (do nothing).
    duplicate_of: int | None = None
    duplicate_action: str | None = None


@dataclass
class ErrorRow:
    row_index: int
    raw_data: dict[str, str]
    error_message: str  # Japanese error message
    error_column: str | None = None  # column that caused the error (for highlighting)


@dataclass
class ValidationResult:
    valid_rows: list[ValidatedRow] = field(default_factory=list)
    error_rows: list[ErrorRow] = field(default_factory=list)
    duplicate_rows: list[ValidatedRow] = field(default_factory=list)


class ValidationPipeline:
    def __init__(
        self, db_manager: DatabaseManager, mapping: ColumnMapping, source_path: str
    ):
        self.db = db_manager
        self.mapping = mapping
        self.source_path = source_path

    def validate(self, df: pd.DataFrame, progress_callback=None) -> ValidationResult:
        """Run full validation pipeline on a DataFrame."""
        result = ValidationResult()
        total = len(df)

        # In-file dedup tracker — (name, death_date_iso) → first row index seen
        seen_in_file: dict[tuple[str, str], int] = {}

        for idx, row in df.iterrows():
            if progress_callback:
                try:
                    progress_callback(int(idx), total)
                except Exception:
                    # Never let a UI signal slot exception abort validation
                    pass

            row_index = int(idx) if isinstance(idx, (int, float)) else 0
            raw = {col: str(row.get(col, "")) for col in df.columns}

            # Stage 1: Pre-validation - check required fields
            name = self._extract_name(row)
            if not name:
                result.error_rows.append(
                    ErrorRow(row_index, raw, "氏名が空です", self.mapping.name_col)
                )
                continue

            # Stage 2: Date parsing
            date_err_col = (
                self.mapping.death_date_col
                if not self.mapping.uses_split_date
                else (self.mapping.year_col or self.mapping.era_col)
            )
            try:
                parsed = self._parse_death_date(row)
            except DateValidationError as e:
                result.error_rows.append(ErrorRow(row_index, raw, str(e), date_err_col))
                continue
            except Exception as e:
                result.error_rows.append(
                    ErrorRow(row_index, raw, f"日付の解析中に予期しないエラー: {e}", date_err_col)
                )
                continue

            # Stage 3: Business rules
            death_date = parsed.date
            if death_date > datetime.date.today():
                result.error_rows.append(
                    ErrorRow(
                        row_index,
                        raw,
                        f"没年月日が未来の日付です: {parsed.era_display}",
                        date_err_col,
                    )
                )
                continue
            if death_date.year < 1868:
                result.error_rows.append(
                    ErrorRow(
                        row_index,
                        raw,
                        f"没年月日が明治以前です: {parsed.era_display}",
                        date_err_col,
                    )
                )
                continue

            # Stage 4: Build attributes (uses column_remap selections)
            try:
                attributes = self._build_attributes(row)
            except Exception as e:
                result.error_rows.append(
                    ErrorRow(row_index, raw, f"属性の構築中にエラー: {e}")
                )
                continue

            iso = death_date.isoformat()

            # Stage 5a: In-file duplicate detection
            key = (name, iso)
            if key in seen_in_file:
                result.error_rows.append(
                    ErrorRow(
                        row_index,
                        raw,
                        f"ファイル内で重複しています（同一データの最初の出現: {seen_in_file[key] + 2}行目）",
                    )
                )
                continue
            seen_in_file[key] = row_index

            validated = ValidatedRow(
                row_index=row_index,
                name=name,
                death_date=iso,
                era_display=parsed.era_display,
                attributes=attributes,
                source_file_path=self.source_path,
            )

            # Stage 5b: DB duplicate detection
            try:
                duplicates = self.db.find_duplicates(name, iso)
            except Exception as e:
                result.error_rows.append(
                    ErrorRow(row_index, raw, f"DB重複チェックエラー: {e}")
                )
                continue
            if duplicates:
                validated.duplicate_of = duplicates[0].id
                result.duplicate_rows.append(validated)
            else:
                result.valid_rows.append(validated)

        if progress_callback:
            try:
                progress_callback(total, total)
            except Exception:
                pass

        return result

    def validate_one(self, raw: dict[str, str], row_index: int) -> ValidatedRow | ErrorRow:
        """Re-validate a single row supplied as a plain dict (used after the
        user edits an erroneous row in the in-UI editor)."""
        row = pd.Series({k: (v if v is not None else "") for k, v in raw.items()})

        name = self._extract_name(row)
        if not name:
            return ErrorRow(row_index, raw, "氏名が空です", self.mapping.name_col)

        date_err_col = (
            self.mapping.death_date_col
            if not self.mapping.uses_split_date
            else (self.mapping.year_col or self.mapping.era_col)
        )
        try:
            parsed = self._parse_death_date(row)
        except DateValidationError as e:
            return ErrorRow(row_index, raw, str(e), date_err_col)
        except Exception as e:
            return ErrorRow(row_index, raw, f"日付の解析中に予期しないエラー: {e}", date_err_col)

        if parsed.date > datetime.date.today():
            return ErrorRow(
                row_index, raw, f"没年月日が未来の日付です: {parsed.era_display}", date_err_col
            )
        if parsed.date.year < 1868:
            return ErrorRow(
                row_index, raw, f"没年月日が明治以前です: {parsed.era_display}", date_err_col
            )

        try:
            attributes = self._build_attributes(row)
        except Exception as e:
            return ErrorRow(row_index, raw, f"属性の構築中にエラー: {e}")

        iso = parsed.date.isoformat()
        validated = ValidatedRow(
            row_index=row_index,
            name=name,
            death_date=iso,
            era_display=parsed.era_display,
            attributes=attributes,
            source_file_path=self.source_path,
        )
        try:
            duplicates = self.db.find_duplicates(name, iso)
        except Exception as e:
            return ErrorRow(row_index, raw, f"DB重複チェックエラー: {e}")
        if duplicates:
            validated.duplicate_of = duplicates[0].id
        return validated

    def _extract_name(self, row: pd.Series) -> str | None:
        col = self.mapping.name_col
        if col is None:
            return None
        val = str(row.get(col, "")).strip()
        return val if val else None

    def _parse_death_date(self, row: pd.Series):
        if self.mapping.uses_split_date:
            year_val = row.get(self.mapping.year_col, "")
            month_val = row.get(self.mapping.month_col, "")
            day_val = row.get(self.mapping.day_col, "")
            era_val = (
                row.get(self.mapping.era_col, "") if self.mapping.era_col else None
            )
            return parse_date_from_columns(year_val, month_val, day_val, era_val)
        else:
            col = self.mapping.death_date_col
            if col is None:
                raise DateValidationError("没年月日の列が設定されていません")
            val = str(row.get(col, "")).strip()
            if not val:
                raise DateValidationError("没年月日が空です")
            return parse_date(val)

    def _build_attributes(self, row: pd.Series) -> dict[str, str]:
        """Build dynamic attributes from extra columns + buddhist name.

        Applies column_remap to rename Excel columns to existing DB column names.
        If multiple source columns end up writing to the same DB column we keep
        the first non-empty value rather than letting the later one silently
        clobber the earlier — the import-window step 3 conflict check should
        prevent this anyway.
        """
        remap = self.mapping.column_remap
        attrs: dict[str, str] = {}

        def _set(col: str, value: str):
            target = remap.get(col, col)
            if not target:
                return
            if target in attrs and attrs[target]:
                return  # keep first non-empty value
            attrs[target] = value

        if self.mapping.buddhist_name_col:
            val = str(row.get(self.mapping.buddhist_name_col, "")).strip()
            if val:
                _set(self.mapping.buddhist_name_col, val)

        for col in self.mapping.extra_cols:
            val = str(row.get(col, "")).strip()
            if val:
                _set(col, val)
        return attrs


def _get_suggestion(error_message: str) -> str:
    """Return a human-readable fix hint for an error message."""
    msg = error_message
    if "氏名が空" in msg:
        return "氏名（必須）を入力してください。"
    if "没年月日が空" in msg:
        return "没年月日（必須）を入力してください。\n例: 令和元年5月1日 / 2019-05-01"
    if "未来の日付" in msg:
        return f"{msg}\n→ 過去の日付を入力してください。"
    if "明治以前" in msg:
        return f"{msg}\n→ 対象は明治元年（1868年）以降のみです。"
    if "ファイル内で重複" in msg:
        return f"{msg}\n→ 重複行を削除または修正してください。"
    if "日付の解析" in msg or "予期しないエラー" in msg:
        return "日付の形式を確認してください。\n正しい形式: 令和○年○月○日 / 平成○年○月○日 / YYYY-MM-DD"
    # DateValidationError messages from the parser — append format hint
    return f"{msg}\n→ 正しい形式: 令和○年○月○日 / 平成○年○月○日 / YYYY-MM-DD"


def create_error_annotated_copy(
    error_rows: list[ErrorRow],
    original_path: Path,
    sheet_name: str | None = None,
) -> "Path | None":
    """Create an annotated copy of the original file with error rows highlighted.

    For xlsx/xlsm files the original sheet is copied and annotated in-place.
    For other formats (xls, csv) a new xlsx is built from the raw row data.
    Returns the path to the created file, or None on failure.
    """
    if not error_rows:
        return None

    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font as OxFont

        suffix = original_path.suffix.lower()
        out_path = original_path.parent / f"{original_path.stem}_エラー.xlsx"

        pink_fill = PatternFill(start_color="FFD9D9", end_color="FFD9D9", fill_type="solid")
        dark_fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
        header_font = OxFont(bold=True, color="CC0000")

        if suffix in (".xlsx", ".xlsm"):
            wb = openpyxl.load_workbook(original_path)
            ws = (
                wb[sheet_name]
                if sheet_name and sheet_name in wb.sheetnames
                else wb.active
            )

            # Map header cell values → column indices
            header_map: dict[str, int] = {
                str(cell.value): cell.column
                for cell in ws[1]
                if cell.value is not None
            }

            suggestion_col = ws.max_column + 1
            hdr = ws.cell(row=1, column=suggestion_col, value="エラー内容・修正提案")
            hdr.font = header_font

            for err in error_rows:
                excel_row = err.row_index + 2  # 0-based pandas + header row + 1-based
                if excel_row > ws.max_row:
                    continue
                for cell in ws[excel_row]:
                    cell.fill = pink_fill
                if err.error_column and err.error_column in header_map:
                    ws.cell(row=excel_row, column=header_map[err.error_column]).fill = dark_fill
                ws.cell(
                    row=excel_row,
                    column=suggestion_col,
                    value=_get_suggestion(err.error_message),
                )
        else:
            # xls / csv: rebuild as a new xlsx with just the error rows
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "エラー行"

            cols = list(error_rows[0].raw_data.keys())
            suggestion_col = len(cols) + 1
            for j, col in enumerate(cols, 1):
                ws.cell(row=1, column=j, value=col)
            hdr = ws.cell(row=1, column=suggestion_col, value="エラー内容・修正提案")
            hdr.font = header_font

            for i, err in enumerate(error_rows, 2):
                for j, col in enumerate(cols, 1):
                    cell = ws.cell(row=i, column=j, value=err.raw_data.get(col, ""))
                    cell.fill = (
                        dark_fill if col == err.error_column else pink_fill
                    )
                ws.cell(
                    row=i,
                    column=suggestion_col,
                    value=_get_suggestion(err.error_message),
                )

        wb.save(out_path)
        return out_path
    except Exception:
        return None


def export_error_rows(error_rows: list[ErrorRow], output_path: Path) -> None:
    """Export error rows as an Excel file for offline correction.

    Cells that start with characters Excel interprets as a formula are prefixed
    with an apostrophe to prevent formula-injection problems when the user
    re-opens the file.
    """
    if not error_rows:
        return
    data = []
    for er in error_rows:
        row_data = {k: _sanitize_for_excel(str(v)) for k, v in er.raw_data.items()}
        row_data["エラー内容"] = _sanitize_for_excel(er.error_message)
        row_data["元の行番号"] = er.row_index + 2  # +2 for header + 0-index
        data.append(row_data)
    df = pd.DataFrame(data)
    df.to_excel(output_path, index=False, engine="openpyxl")


def import_validated_rows(db: DatabaseManager, rows: list[ValidatedRow]):
    """Import validated rows into the database.

    Honours each row's duplicate_action:
      - None / "insert" → insert as a new record
      - "overwrite"     → update the existing row pointed to by duplicate_of
      - "skip"          → do nothing for this row

    Returns a BatchResult-like dict with inserted/updated/skipped counts and
    per-row failures.
    """
    inserted_records = []
    updated = 0
    skipped = 0
    failures: list[tuple[int, str]] = []

    for row in rows:
        action = row.duplicate_action
        if action == "skip":
            skipped += 1
            continue
        if action == "overwrite" and row.duplicate_of is not None:
            try:
                db.update_person(
                    row.duplicate_of,
                    name=row.name,
                    death_date=row.death_date,
                    attributes=row.attributes,
                    merge_attributes=True,
                )
                updated += 1
            except Exception as e:
                failures.append((row.row_index, str(e)))
            continue
        inserted_records.append(
            {
                "name": row.name,
                "death_date": row.death_date,
                "source_file_path": row.source_file_path,
                "attributes": row.attributes,
            }
        )

    insert_result = db.add_persons_batch(inserted_records) if inserted_records else None
    inserted = insert_result.inserted if insert_result else 0
    if insert_result:
        failures.extend(insert_result.failures)
    return {
        "inserted": inserted,
        "updated": updated,
        "skipped": skipped,
        "failures": failures,
    }
