"""Word document generator for nenki anniversary tables.

Generates .docx files with vertical Japanese text (縦書き / tategaki),
landscape A4 orientation, and proper Japanese formatting.
Field values are padded to fixed widths so that columns align across entries.
Optionally converts to PDF via docx2pdf.
"""

from pathlib import Path
import unicodedata

from docx import Document
from docx.shared import Pt, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

# Register extra namespaces needed for text boxes
nsmap['wps'] = 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'
nsmap['wp14'] = 'http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing'
nsmap['mc'] = 'http://schemas.openxmlformats.org/markup-compatibility/2006'


def _display_width(text: str) -> int:
    """Calculate display width treating fullwidth/CJK chars as 2, others as 1."""
    w = 0
    for ch in text:
        eaw = unicodedata.east_asian_width(ch)
        w += 2 if eaw in ("F", "W", "A") else 1
    return w


def _pad_to_width(text: str, target_width: int) -> str:
    """Pad text with full-width spaces to reach target display width."""
    current = _display_width(text)
    # Each full-width space has display width 2
    needed = (target_width - current) // 2
    if needed > 0:
        return text + "\u3000" * needed
    return text


class WordGenerator:
    """Generate Word documents for nenki anniversary listings with vertical text."""

    FONT_NAME = "MS Mincho"
    # Full-width space used as field separator
    FIELD_SEP = "\u3000"

    def create_combined_document(
        self,
        sorted_data: list,
        title: str,
        output_path: Path,
        field_names: list[str] | None = None,
        single_column: bool = True,
    ):
        """Create a combined nenki document with vertical Japanese text.

        Args:
            sorted_data: List of (key, entries).
                key: "年忌名|years_offset"
                entries: list of field-value lists (ordered by field_names)
            title: Document title (e.g. "年忌表 - 令和八年の年忌法要")
            output_path: Path to save the .docx file.
            field_names: List of field names corresponding to each entry's values.
            single_column: True for single column, False for dual column layout.
        """
        doc = Document()

        section = doc.sections[0]
        self._setup_section(section)

        # Set default font
        style = doc.styles["Normal"]
        style.font.name = self.FONT_NAME
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = rPr.makeelement(qn("w:rFonts"), {})
            rPr.insert(0, rFonts)
        rFonts.set(qn("w:eastAsia"), self.FONT_NAME)

        # Compute global field widths across ALL groups for consistent alignment
        all_entries = []
        for _, people_data in sorted_data:
            all_entries.extend(people_data)
        field_widths = self._compute_field_widths(all_entries, field_names)

        if single_column:
            self._add_title(doc, title)
            self._build_single_column_content(doc, sorted_data, field_names, field_widths)
        else:
            # Two columns on the section
            sectPr = section._sectPr
            cols = OxmlElement("w:cols")
            cols.set(qn("w:num"), "2")
            cols.set(qn("w:space"), "720")
            sectPr.append(cols)

            # Floating text box for title — independent of column layout
            self._add_textbox_title(doc, title)

            # Data flows naturally into 2 columns
            self._build_dual_column_content(doc, sorted_data, field_widths)

        doc.save(str(output_path))

        # Auto-convert to PDF
        self._convert_to_pdf(output_path)

    def _add_title(self, doc: Document, title: str):
        """Add full-width centered title (single column mode)."""
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(title)
        title_run.font.size = Pt(36)
        title_run.font.bold = True
        title_run.font.name = self.FONT_NAME
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_para.paragraph_format.space_after = Pt(4)

    def _add_textbox_title(self, doc: Document, title: str):
        """Add title as a floating tategaki text box spanning full page width.

        The text box is anchored at the top of the margin area with
        'top and bottom' wrapping so the 2-column data flows below it.
        """
        para = doc.add_paragraph()
        run = para.add_run()

        # --- Build the DrawingML anchor with a WordprocessingShape text box ---

        # Usable dimensions (page minus margins)
        # Landscape A4: 11.69" x 8.27", margins 0.5" each side
        usable_w = int(10.69 * 914400)  # EMU (width in landscape = physical width - margins)
        title_h = int(0.75 * 914400)     # EMU — text box height (~0.75 inch)

        drawing = OxmlElement('w:drawing')

        # wp:anchor — floating positioning
        anchor = OxmlElement('wp:anchor')
        for attr, val in [
            ('distT', '0'), ('distB', '0'), ('distL', '0'), ('distR', '0'),
            ('simplePos', '0'), ('relativeHeight', '251659264'),
            ('behindDoc', '0'), ('locked', '0'),
            ('layoutInCell', '1'), ('allowOverlap', '1'),
        ]:
            anchor.set(attr, val)

        # Simple position (required but unused)
        simplePos = OxmlElement('wp:simplePos')
        simplePos.set('x', '0')
        simplePos.set('y', '0')
        anchor.append(simplePos)

        # Horizontal: centered relative to margin
        posH = OxmlElement('wp:positionH')
        posH.set('relativeFrom', 'margin')
        align_h = OxmlElement('wp:align')
        align_h.text = 'center'
        posH.append(align_h)
        anchor.append(posH)

        # Vertical: top of margin
        posV = OxmlElement('wp:positionV')
        posV.set('relativeFrom', 'margin')
        offset_v = OxmlElement('wp:posOffset')
        offset_v.text = '0'
        posV.append(offset_v)
        anchor.append(posV)

        # Extent (size)
        extent = OxmlElement('wp:extent')
        extent.set('cx', str(usable_w))
        extent.set('cy', str(title_h))
        anchor.append(extent)

        # Effect extent
        effectExtent = OxmlElement('wp:effectExtent')
        for attr in ('l', 't', 'r', 'b'):
            effectExtent.set(attr, '0')
        anchor.append(effectExtent)

        # Wrap: top and bottom — columns flow below the text box
        anchor.append(OxmlElement('wp:wrapTopAndBottom'))

        # Document properties
        docPr = OxmlElement('wp:docPr')
        docPr.set('id', '1')
        docPr.set('name', 'Title Text Box')
        anchor.append(docPr)

        # Graphic frame
        graphic = OxmlElement('a:graphic')
        graphicData = OxmlElement('a:graphicData')
        graphicData.set('uri', 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape')
        graphic.append(graphicData)

        # Word processing shape
        wsp = OxmlElement('wps:wsp')

        # Shape properties — non-visual
        cNvSpPr = OxmlElement('wps:cNvSpPr')
        cNvSpPr.set('txBox', '1')
        wsp.append(cNvSpPr)

        # Shape properties — geometry, no fill, no border
        spPr = OxmlElement('wps:spPr')

        xfrm = OxmlElement('a:xfrm')
        off = OxmlElement('a:off')
        off.set('x', '0')
        off.set('y', '0')
        xfrm.append(off)
        ext = OxmlElement('a:ext')
        ext.set('cx', str(usable_w))
        ext.set('cy', str(title_h))
        xfrm.append(ext)
        spPr.append(xfrm)

        prstGeom = OxmlElement('a:prstGeom')
        prstGeom.set('prst', 'rect')
        prstGeom.append(OxmlElement('a:avLst'))
        spPr.append(prstGeom)

        spPr.append(OxmlElement('a:noFill'))

        ln = OxmlElement('a:ln')
        ln.append(OxmlElement('a:noFill'))
        spPr.append(ln)

        wsp.append(spPr)

        # Text box content
        txbx = OxmlElement('wps:txbx')
        txbxContent = OxmlElement('w:txbxContent')

        # Title paragraph inside text box
        tp = OxmlElement('w:p')
        tpPr = OxmlElement('w:pPr')
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), 'center')
        tpPr.append(jc)
        tp.append(tpPr)

        tr = OxmlElement('w:r')
        trPr = OxmlElement('w:rPr')
        # Bold
        trPr.append(OxmlElement('w:b'))
        # Font size: 36pt = 72 half-points
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '72')
        trPr.append(sz)
        szCs = OxmlElement('w:szCs')
        szCs.set(qn('w:val'), '72')
        trPr.append(szCs)
        # Font name
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), self.FONT_NAME)
        rFonts.set(qn('w:eastAsia'), self.FONT_NAME)
        rFonts.set(qn('w:hAnsi'), self.FONT_NAME)
        trPr.append(rFonts)
        tr.append(trPr)

        tt = OxmlElement('w:t')
        tt.set(qn('xml:space'), 'preserve')
        tt.text = title
        tr.append(tt)
        tp.append(tr)

        txbxContent.append(tp)
        txbx.append(txbxContent)
        wsp.append(txbx)

        # Body properties — tategaki vertical text, centered anchor
        bodyPr = OxmlElement('wps:bodyPr')
        bodyPr.set('vert', 'eaVert')
        bodyPr.set('wrap', 'square')
        bodyPr.set('lIns', '91440')
        bodyPr.set('tIns', '45720')
        bodyPr.set('rIns', '91440')
        bodyPr.set('bIns', '45720')
        bodyPr.set('anchor', 'ctr')
        wsp.append(bodyPr)

        graphicData.append(wsp)
        anchor.append(graphic)
        drawing.append(anchor)
        run._element.append(drawing)

    def _setup_section(self, section):
        """Configure a section with landscape A4, tategaki, and margins."""
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

        # Set vertical text direction (tategaki)
        sectPr = section._sectPr
        textDirection = OxmlElement("w:textDirection")
        textDirection.set(qn("w:val"), "tbRl")
        sectPr.append(textDirection)

    def _compute_field_widths(self, all_entries: list, field_names: list[str] | None) -> list[int]:
        """Compute the max display width for each field position across all entries."""
        if not all_entries:
            return []
        num_fields = len(all_entries[0]) if all_entries else 0
        widths = [0] * num_fields

        for entry in all_entries:
            for i, val in enumerate(entry):
                w = _display_width(str(val) if val else "")
                if w > widths[i]:
                    widths[i] = w

        # Round up each width to even number (full-width space alignment)
        widths = [w + (w % 2) for w in widths]
        return widths

    def _format_aligned_entry(self, entry: list, field_widths: list[int]) -> str:
        """Format entry fields padded to fixed widths for alignment."""
        parts = []
        for i, val in enumerate(entry):
            text = str(val) if val else ""
            if i < len(field_widths):
                text = _pad_to_width(text, field_widths[i])
            parts.append(text)
        return self.FIELD_SEP.join(parts)

    def _build_single_column_content(
        self, doc: Document, sorted_data: list,
        field_names: list[str] | None, field_widths: list[int],
    ):
        """Single column layout: each nenki group with subtitle + entries."""
        for key, people_data in sorted_data:
            nenki_name = key.split("|")[0]

            subtitle = doc.add_paragraph()
            subtitle_run = subtitle.add_run(nenki_name)
            subtitle_run.font.size = Pt(28)
            subtitle_run.font.bold = True
            subtitle_run.font.name = self.FONT_NAME
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle.paragraph_format.space_after = Pt(12)

            for entry in people_data:
                text = self.FIELD_SEP + self._format_aligned_entry(entry, field_widths)
                para = doc.add_paragraph()
                run = para.add_run(text)
                run.font.size = Pt(18)
                run.font.name = self.FONT_NAME
                para.paragraph_format.left_indent = Inches(1.5)
                para.paragraph_format.space_after = Pt(8)

            # Space between groups
            doc.add_paragraph()

    def _build_dual_column_content(
        self, doc: Document, sorted_data: list, field_widths: list[int],
    ):
        """Dual column layout: data flows into Word's native 2 columns."""
        for key, people_data in sorted_data:
            nenki_name = key.split("|")[0]

            subtitle = doc.add_paragraph()
            subtitle_run = subtitle.add_run(nenki_name)
            subtitle_run.font.size = Pt(16)
            subtitle_run.font.bold = True
            subtitle_run.font.name = self.FONT_NAME
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle.paragraph_format.space_after = Pt(8)

            for entry in people_data:
                text = self.FIELD_SEP + self._format_aligned_entry(entry, field_widths)
                para = doc.add_paragraph()
                run = para.add_run(text)
                run.font.size = Pt(11)
                run.font.name = self.FONT_NAME
                para.paragraph_format.left_indent = Inches(0.5)
                para.paragraph_format.space_after = Pt(4)

            # Space between groups
            doc.add_paragraph()

    def _convert_to_pdf(self, docx_path: Path):
        """Try to convert .docx to .pdf using docx2pdf."""
        try:
            from docx2pdf import convert
            pdf_path = docx_path.with_suffix(".pdf")
            convert(str(docx_path), str(pdf_path))
        except Exception:
            pass
